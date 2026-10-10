"""
Excel导出模块 - 用于生成详细的性能测试报告
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, LineChart, Reference
import json
from datetime import datetime
def export_to_excel(all_results, model_name, use_long_context, output_file='benchmark_results.xlsx'):
    """
    将性能测试结果导出为Excel文件
    
    Args:
        all_results: 所有测试结果列表
        model_name: 模型名称
        use_long_context: 是否使用长文本模式
        output_file: 输出文件名
    """
    wb = Workbook()
    
    # 删除默认的Sheet
    if 'Sheet' in wb.sheetnames:
        wb.remove(wb['Sheet'])
    
    # 创建多个工作表
    create_summary_sheet(wb, all_results, model_name, use_long_context)
    create_detailed_sheet(wb, all_results)
    create_latency_analysis_sheet(wb, all_results)
    create_throughput_analysis_sheet(wb, all_results)
    
    # 保存文件
    wb.save(output_file)
    print(f"✅ Excel报告已保存至: {output_file}")
    return output_file
def create_summary_sheet(wb, all_results, model_name, use_long_context):
    """创建汇总工作表"""
    ws = wb.create_sheet("测试汇总", 0)
    
    # 设置列宽
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 30
    
    # 标题样式
    title_font = Font(name='微软雅黑', size=16, bold=True, color='FFFFFF')
    title_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='5B9BD5', end_color='5B9BD5', fill_type='solid')
    
    # 添加标题
    ws['A1'] = '性能测试汇总报告'
    ws['A1'].font = title_font
    ws['A1'].fill = title_fill
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.merge_cells('A1:B1')
    ws.row_dimensions[1].height = 30
    
    # 基本信息
    row = 3
    ws[f'A{row}'] = '基本信息'
    ws[f'A{row}'].font = header_font
    ws[f'A{row}'].fill = header_fill
    ws.merge_cells(f'A{row}:B{row}')
    
    # 计算汇总数据
    total_tokens = sum(r.get('total_output_tokens', 0) for r in all_results)
    total_time = sum(r.get('total_time', 0) for r in all_results)
    total_requests = sum(r.get('total_requests', 0) for r in all_results)
    successful_requests = sum(r.get('successful_requests', 0) for r in all_results)
    
    info_data = [
        ['测试时间', datetime.now().strftime('%Y-%m-%d %H:%M:%S')],
        ['测试模型', model_name],
        ['长文本模式', '是' if use_long_context else '否'],
        ['测试配置数', len(all_results)],
        ['总请求数', total_requests],
        ['成功请求数', successful_requests],
        ['总成功率', f"{(successful_requests/total_requests*100):.2f}%" if total_requests > 0 else 'N/A'],
        ['总生成Token数', f"{total_tokens:,}"],
        ['总测试时间(秒)', f"{total_time:.2f}"],
        ['平均Token生成速率', f"{total_tokens/total_time:.2f} tokens/sec" if total_time > 0 else 'N/A'],
    ]
    
    row += 1
    for item in info_data:
        ws[f'A{row}'] = item[0]
        ws[f'B{row}'] = item[1]
        ws[f'A{row}'].font = Font(name='微软雅黑', bold=True)
        row += 1
    
    # 性能指标汇总表
    row += 2
    ws[f'A{row}'] = '性能指标汇总'
    ws[f'A{row}'].font = header_font
    ws[f'A{row}'].fill = header_fill
    ws.merge_cells(f'A{row}:H{row}')
    
    # 表头
    row += 1
    headers = ['并发数', '请求数', 'RPS', '平均延迟(秒)', 'P99延迟(秒)', '平均TPS', '首Token延迟(秒)', '成功率']
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col)
        cell.value = header
        cell.font = Font(name='微软雅黑', bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='70AD47', end_color='70AD47', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.column_dimensions[get_column_letter(col)].width = 15
    
    # 数据行
    for result in all_results:
        row += 1
        concurrency = result.get('concurrency', 0)
        total_req = result.get('total_requests', 0)
        rps = result.get('requests_per_second', 0)
        avg_latency = result.get('latency', {}).get('average', 0)
        p99_latency = result.get('latency', {}).get('p99', 0)
        avg_tps = result.get('tokens_per_second', {}).get('average', 0)
        avg_ttft = result.get('time_to_first_token', {}).get('average', 0)
        success_rate = (result.get('successful_requests', 0) / total_req * 100) if total_req > 0 else 0
        
        data = [
            concurrency,
            total_req,
            round(rps, 2) if rps is not None else 0,
            round(avg_latency, 3) if avg_latency is not None else 0,
            round(p99_latency, 3) if p99_latency is not None else 0,
            round(avg_tps, 2) if avg_tps is not None else 0,
            round(avg_ttft, 3) if avg_ttft is not None else 0,
            f"{success_rate:.1f}%"
        ]
        
        for col, value in enumerate(data, start=1):
            cell = ws.cell(row=row, column=col)
            cell.value = value
            cell.alignment = Alignment(horizontal='center', vertical='center')
            
            # 根据成功率设置颜色
            if col == 8:  # 成功率列
                if success_rate >= 95:
                    cell.fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
                elif success_rate >= 80:
                    cell.fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
                else:
                    cell.fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
    
    # 添加边框
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    for row_cells in ws.iter_rows(min_row=row-len(all_results), max_row=row, min_col=1, max_col=8):
        for cell in row_cells:
            cell.border = thin_border
def create_detailed_sheet(wb, all_results):
    """创建详细数据工作表"""
    ws = wb.create_sheet("详细数据")
    
    # 表头
    headers = [
        '并发数', '总请求数', '成功请求数', '失败请求数', '成功率(%)',
        'RPS', '平均延迟(秒)', '最小延迟(秒)', '最大延迟(秒)', 
        'P50延迟(秒)', 'P95延迟(秒)', 'P99延迟(秒)',
        '平均TPS', '最小TPS', '最大TPS',
        '平均首Token延迟(秒)', '最小首Token延迟(秒)', '最大首Token延迟(秒)',
        '总输入Token数', '总输出Token数', '总测试时间(秒)'
    ]
    
    # 设置表头样式
    header_font = Font(name='微软雅黑', size=10, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = 14
    
    ws.row_dimensions[1].height = 30
    
    # 填充数据
    for idx, result in enumerate(all_results, start=2):
        total_req = result.get('total_requests', 0)
        success_req = result.get('successful_requests', 0)
        failed_req = total_req - success_req
        success_rate = (success_req / total_req * 100) if total_req > 0 else 0
        
        latency = result.get('latency', {})
        tps = result.get('tokens_per_second', {})
        ttft = result.get('time_to_first_token', {})
        
        data = [
            result.get('concurrency', 0),
            total_req,
            success_req,
            failed_req,
            round(success_rate, 2),
            round(result.get('requests_per_second', 0) or 0, 2),
            round(latency.get('average', 0) or 0, 3),
            round(latency.get('min', 0) or 0, 3),
            round(latency.get('max', 0) or 0, 3),
            round(latency.get('p50', 0) or 0, 3),
            round(latency.get('p95', 0) or 0, 3),
            round(latency.get('p99', 0) or 0, 3),
            round(tps.get('average', 0) or 0, 2),
            round(tps.get('min', 0) or 0, 2),
            round(tps.get('max', 0) or 0, 2),
            round(ttft.get('average', 0) or 0, 3),
            round(ttft.get('min', 0) or 0, 3),
            round(ttft.get('max', 0) or 0, 3),
            result.get('total_input_tokens', 0),
            result.get('total_output_tokens', 0),
            round(result.get('total_time', 0) or 0, 2)
        ]
        
        for col, value in enumerate(data, start=1):
            cell = ws.cell(row=idx, column=col)
            cell.value = value
            cell.alignment = Alignment(horizontal='center', vertical='center')
            
            # 成功率着色
            if col == 5:
                if value >= 95:
                    cell.fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
                elif value >= 80:
                    cell.fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
                else:
                    cell.fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
def create_latency_analysis_sheet(wb, all_results):
    """创建延迟分析工作表（带图表）"""
    ws = wb.create_sheet("延迟分析")
    
    # 表头
    headers = ['并发数', '平均延迟(秒)', 'P50延迟(秒)', 'P95延迟(秒)', 'P99延迟(秒)', '最大延迟(秒)']
    header_font = Font(name='微软雅黑', size=10, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='ED7D31', end_color='ED7D31', fill_type='solid')
    
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.column_dimensions[get_column_letter(col)].width = 15
    
    # 填充数据
    for idx, result in enumerate(all_results, start=2):
        latency = result.get('latency', {})
        data = [
            result.get('concurrency', 0),
            round(latency.get('average', 0) or 0, 3),
            round(latency.get('p50', 0) or 0, 3),
            round(latency.get('p95', 0) or 0, 3),
            round(latency.get('p99', 0) or 0, 3),
            round(latency.get('max', 0) or 0, 3)
        ]
        
        for col, value in enumerate(data, start=1):
            ws.cell(row=idx, column=col, value=value).alignment = Alignment(horizontal='center')
    
    # 创建折线图
    chart = LineChart()
    chart.title = "延迟随并发数变化趋势"
    chart.style = 10
    chart.y_axis.title = '延迟(秒)'
    chart.x_axis.title = '并发数'
    chart.height = 12
    chart.width = 20
    
    # 添加数据系列
    data = Reference(ws, min_col=2, min_row=1, max_row=len(all_results)+1, max_col=6)
    cats = Reference(ws, min_col=1, min_row=2, max_row=len(all_results)+1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    
    ws.add_chart(chart, "H2")
def create_throughput_analysis_sheet(wb, all_results):
    """创建吞吐量分析工作表（带图表）"""
    ws = wb.create_sheet("吞吐量分析")
    
    # 表头
    headers = ['并发数', 'RPS', '平均TPS', '最大TPS', '总输出Token数', '测试时间(秒)']
    header_font = Font(name='微软雅黑', size=10, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='70AD47', end_color='70AD47', fill_type='solid')
    
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.column_dimensions[get_column_letter(col)].width = 16
    
    # 填充数据
    for idx, result in enumerate(all_results, start=2):
        tps = result.get('tokens_per_second', {})
        data = [
            result.get('concurrency', 0),
            round(result.get('requests_per_second', 0) or 0, 2),
            round(tps.get('average', 0) or 0, 2),
            round(tps.get('max', 0) or 0, 2),
            result.get('total_output_tokens', 0),
            round(result.get('total_time', 0) or 0, 2)
        ]
        
        for col, value in enumerate(data, start=1):
            ws.cell(row=idx, column=col, value=value).alignment = Alignment(horizontal='center')
    
    # 创建柱状图 - RPS
    chart1 = BarChart()
    chart1.type = "col"
    chart1.title = "RPS随并发数变化"
    chart1.y_axis.title = 'RPS (请求/秒)'
    chart1.x_axis.title = '并发数'
    chart1.height = 10
    chart1.width = 18
    
    data = Reference(ws, min_col=2, min_row=1, max_row=len(all_results)+1)
    cats = Reference(ws, min_col=1, min_row=2, max_row=len(all_results)+1)
    chart1.add_data(data, titles_from_data=True)
    chart1.set_categories(cats)
    
    ws.add_chart(chart1, "H2")
    
    # 创建柱状图 - TPS
    chart2 = BarChart()
    chart2.type = "col"
    chart2.title = "平均TPS随并发数变化"
    chart2.y_axis.title = 'TPS (tokens/秒)'
    chart2.x_axis.title = '并发数'
    chart2.height = 10
    chart2.width = 18
    
    data = Reference(ws, min_col=3, min_row=1, max_row=len(all_results)+1)
    chart2.add_data(data, titles_from_data=True)
    chart2.set_categories(cats)
    
    ws.add_chart(chart2, "H18")
if __name__ == "__main__":
    print("✅ Excel导出模块已创建")
    print("📋 包含以下功能:")
    print("  - 测试汇总工作表（基本信息 + 性能指标汇总）")
    print("  - 详细数据工作表（完整的测试数据）")
    print("  - 延迟分析工作表（延迟趋势图表）")
    print("  - 吞吐量分析工作表（RPS和TPS图表）")